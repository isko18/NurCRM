from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.users.models import Branch, Company, User
from apps.warehouse.models import Warehouse, WarehouseProduct, WarehouseProductGroup


YELLOW_RGB = "FFFFFF00"


def clean_value(value):
    if value is None:
        return ""
    return str(value).strip()


def decimal_value(value):
    raw = clean_value(value)
    if not raw:
        return Decimal("0")
    try:
        return Decimal(raw.replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def is_yellow(cell):
    color = cell.fill.fgColor
    return color.type == "rgb" and color.rgb == YELLOW_RGB


class Command(BaseCommand):
    help = "Импорт прайс-листа в warehouse с разделением товаров по желтым группам."

    def add_arguments(self, parser):
        parser.add_argument("file", help="Путь к .xlsx файлу")
        parser.add_argument("--email", required=True, help="Email владельца/пользователя компании")
        parser.add_argument("--warehouse-name", help="Название склада, если у компании несколько складов")
        parser.add_argument("--branch", help="UUID филиала. Если не указан, берется из склада")
        parser.add_argument("--sheet", default=None, help="Название листа. Если не указано, берется активный лист")
        parser.add_argument("--name-col", type=int, default=2, help="Колонка названия товара (1-based)")
        parser.add_argument("--purchase-price-col", type=int, default=3, help="Колонка цены покупки (1-based)")
        parser.add_argument("--price-col", type=int, default=4, help="Колонка цены продажи (1-based)")
        parser.add_argument("--start-row", type=int, default=1, help="Строка начала чтения (1-based)")
        parser.add_argument("--dry-run", action="store_true", help="Только показать, без записи в БД")

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("Не установлен openpyxl") from exc

        user = User.objects.filter(email__iexact=options["email"]).select_related("company").first()
        if not user:
            raise CommandError(f"Пользователь с email не найден: {options['email']}")

        company = user.company or Company.objects.filter(owner=user).first()
        if not company:
            raise CommandError("У найденного пользователя нет компании")

        branch = None
        if options.get("branch"):
            branch = Branch.objects.filter(id=options["branch"], company=company).first()
            if not branch:
                raise CommandError(f"Филиал не найден: {options['branch']}")

        warehouses = Warehouse.objects.filter(company=company)
        if branch:
            warehouses = warehouses.filter(branch=branch)
        if options.get("warehouse_name"):
            warehouses = warehouses.filter(name__iexact=options["warehouse_name"])

        warehouses_count = warehouses.count()
        if warehouses_count == 0:
            raise CommandError("Склад компании не найден")
        if warehouses_count > 1:
            names = ", ".join(clean_value(w.name) or str(w.id) for w in warehouses[:10])
            raise CommandError(f"У компании несколько складов. Укажите --warehouse-name. Склады: {names}")

        warehouse = warehouses.first()
        if branch is None:
            branch = warehouse.branch

        try:
            wb = load_workbook(options["file"], data_only=True)
        except Exception as exc:
            raise CommandError(f"Не удалось прочитать файл: {exc}") from exc

        try:
            ws = wb[options["sheet"]] if options["sheet"] else wb.active
        except KeyError as exc:
            raise CommandError(f"Лист не найден: {options['sheet']}") from exc

        current_group = None
        created_groups = 0
        created_products = 0
        updated_products = 0
        skipped = 0
        errors = []

        @transaction.atomic
        def import_rows():
            nonlocal current_group, created_groups, created_products, updated_products, skipped

            for row_num in range(options["start_row"], ws.max_row + 1):
                name_cell = ws.cell(row_num, options["name_col"])
                name = clean_value(name_cell.value)
                if not name:
                    skipped += 1
                    continue

                row_is_group = any(is_yellow(ws.cell(row_num, col)) for col in range(1, ws.max_column + 1))
                if row_is_group:
                    current_group, created = WarehouseProductGroup.objects.get_or_create(
                        warehouse=warehouse,
                        parent=None,
                        name=name[:128],
                        defaults={"company": company, "branch": branch},
                    )
                    if created:
                        created_groups += 1
                    continue

                purchase_price = decimal_value(ws.cell(row_num, options["purchase_price_col"]).value)
                price = decimal_value(ws.cell(row_num, options["price_col"]).value)
                product = (
                    WarehouseProduct.objects
                    .filter(warehouse=warehouse, company=company, name=name[:255])
                    .first()
                )

                try:
                    if product:
                        product.branch = branch
                        product.product_group = current_group
                        product.purchase_price = purchase_price
                        product.price = price
                        product.full_clean()
                        product.save()
                        updated_products += 1
                    else:
                        product = WarehouseProduct(
                            company=company,
                            branch=branch,
                            warehouse=warehouse,
                            product_group=current_group,
                            name=name[:255],
                            purchase_price=purchase_price,
                            price=price,
                            quantity=Decimal("0"),
                            unit="шт.",
                            status=WarehouseProduct.Status.ACCEPTED,
                        )
                        product.full_clean()
                        product.save()
                        created_products += 1
                except Exception as exc:
                    errors.append(f"Строка {row_num}: {exc}")

            if options["dry_run"]:
                transaction.set_rollback(True)

        import_rows()
        wb.close()

        prefix = "[DRY-RUN] " if options["dry_run"] else ""
        self.stdout.write(
            self.style.SUCCESS(
                f"{prefix}Импорт завершен. "
                f"Групп создано: {created_groups}. "
                f"Товаров создано: {created_products}, обновлено: {updated_products}, пропущено: {skipped}."
            )
        )

        if errors:
            self.stdout.write(self.style.WARNING("Ошибки:"))
            for error in errors[:20]:
                self.stdout.write(f" - {error}")
            if len(errors) > 20:
                self.stdout.write(f"... еще {len(errors) - 20} ошибок")
