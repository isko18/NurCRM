"""
Экспорт всех товаров компании в Excel.

Компания определяется по email владельца (Company.owner.email).

Запуск:
    python export_products.py ormonata@gmail.com
    python export_products.py ormonata@gmail.com --out products.xlsx

Если email не указан — берётся значение DEFAULT_EMAIL ниже.
"""

import os
import sys
import argparse
from decimal import Decimal

import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "core.settings")
django.setup()

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.users.models import Company
from apps.main.models import Product

DEFAULT_EMAIL = "ormonata@gmail.com"

# (Заголовок в Excel, как достать значение из объекта Product)
COLUMNS = [
    ("Код", lambda p: p.code),
    ("Артикул", lambda p: p.article),
    ("Название", lambda p: p.name),
    ("Описание", lambda p: p.description),
    ("Штрихкод", lambda p: p.barcode),
    ("Тип", lambda p: p.get_kind_display()),
    ("Бренд", lambda p: p.brand.name if p.brand else ""),
    ("Категория", lambda p: p.category.name if p.category else ""),
    ("Филиал", lambda p: p.branch.name if p.branch else ""),
    ("Единица", lambda p: p.unit),
    ("Весовой", lambda p: "да" if p.is_weight else "нет"),
    ("Количество/Остаток", lambda p: p.quantity),
    ("Мин. остаток", lambda p: p.minimum_quantity),
    ("Цена закупки", lambda p: p.purchase_price),
    ("Наценка, %", lambda p: p.markup_percent),
    ("Цена продажи", lambda p: p.price),
    ("Оптовая цена", lambda p: p.wholesale_price),
    ("Скидка, %", lambda p: p.discount_percent),
    ("ПЛУ", lambda p: p.plu),
    ("Создан", lambda p: p.created_at.replace(tzinfo=None) if getattr(p, "created_at", None) else ""),
]


def main():
    parser = argparse.ArgumentParser(description="Экспорт товаров компании в Excel")
    parser.add_argument("email", nargs="?", default=DEFAULT_EMAIL,
                        help="Email владельца компании")
    parser.add_argument("--out", default=None, help="Имя выходного .xlsx файла")
    args = parser.parse_args()

    try:
        company = Company.objects.select_related("owner").get(owner__email__iexact=args.email)
    except Company.DoesNotExist:
        sys.exit(f"Компания с владельцем {args.email} не найдена")
    except Company.MultipleObjectsReturned:
        sys.exit(f"Найдено несколько компаний с владельцем {args.email}")

    products = (
        Product.objects
        .filter(company=company)
        .select_related("brand", "category", "branch")
        .order_by("code", "name")
    )

    out_path = args.out or f"products_{company.slug or company.id}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"

    # Заголовки
    headers = [c[0] for c in COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    count = 0
    for p in products.iterator():
        row = []
        for _, getter in COLUMNS:
            val = getter(p)
            if isinstance(val, Decimal):
                val = float(val)
            row.append(val)
        ws.append(row)
        count += 1

    # Автоширина колонок (по содержимому)
    for idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        col = get_column_letter(idx)
        for cell in ws[col]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col].width = min(max_len + 2, 50)

    ws.freeze_panes = "A2"
    wb.save(out_path)

    print(f"Компания: {company.name} (владелец {args.email})")
    print(f"Выгружено товаров: {count}")
    print(f"Файл: {os.path.abspath(out_path)}")


if __name__ == "__main__":
    main()
